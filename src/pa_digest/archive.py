from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Article
from .state import StateStore


HEADERS = [
    "저자",
    "연도",
    "제목",
    "저널명",
    "방법론 (크게)",
    "방법론 (세부)",
    "키워드 1",
    "키워드 2",
    "한글 요약",
    "Abstract",
]


def export_archive(state: StateStore, path: Path) -> None:
    """Write one de-duplicated, sortable workbook from all prepared/sent batches."""
    records: dict[str, Article] = {}
    for batch in state.data.get("batches", {}).values():
        if batch.get("status") not in {"prepared", "sent"}:
            continue
        for record in batch.get("items", []):
            article = Article.model_validate(record)
            # A normal sent batch wins over an explicit resend of the same paper.
            records.setdefault(article.stable_id, article)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "논문 요약"
    sheet.append(HEADERS)
    header_fill = PatternFill("solid", fgColor="302F2B")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:J1"

    for article in sorted(records.values(), key=lambda item: (item.publication_date, item.title), reverse=True):
        constructs = article.constructs + ["", ""]
        sheet.append([
            "; ".join(article.authors),
            article.publication_date.year,
            article.title,
            article.journal,
            article.method or "",
            article.method_detail or "",
            constructs[0],
            constructs[1],
            article.summary_ko or "",
            article.abstract or "",
        ])
        title_cell = sheet.cell(sheet.max_row, 3)
        title_cell.hyperlink = article.url
        title_cell.style = "Hyperlink"

    widths = [30, 10, 52, 42, 18, 20, 22, 22, 60, 90]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.row_dimensions[1].height = 24
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
